#!/usr/bin/env python3
"""Generate RaidInspector_BiS.lua from the "[WotLK] BiS Lists for 3.3.5a end game" workbook.

The workbook has one sheet per spec. Each row is one gear slot: column A is the
slot, the best-in-slot item sits in the first hyperlinked cell (the link target
carries the item id, e.g. ...cavernoftime.com/item=50640) with its iLvl in the
nearest numeric cell to its left; any further hyperlinked cells on the row are
alternatives. A "Navigation Panel" sheet groups the specs by role.

Usage:
    python tools/build_bis_data.py [workbook.xlsx] [--output RaidInspector_BiS.lua]

Requires openpyxl. The output is deterministic so it diffs cleanly in git.
"""
import argparse
import datetime
import os
import re
import sys

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

DEFAULT_WORKBOOK = r"C:\Users\scava\Desktop\bis list\[WotLK] BiS Lists for 3.3.5a end game.xlsx"
ITEM_LINK = re.compile(r"item=(\d+)")
NAV_SHEET = "Navigation Panel"

# Sheet-name keyword -> WoW class token (as used by RAID_CLASS_COLORS).
CLASS_KEYWORDS = [
    ("DK", "DEATHKNIGHT"),
    ("Warrior", "WARRIOR"),
    ("Paladin", "PALADIN"),
    ("Rogue", "ROGUE"),
    ("Druid", "DRUID"),
    ("Feral Cat", "DRUID"),
    ("Bear", "DRUID"),
    ("Hunter", "HUNTER"),
    ("Mage", "MAGE"),
    ("Warlock", "WARLOCK"),
    ("Priest", "PRIEST"),
    ("Shaman", "SHAMAN"),
    ("Spellhancement", "SHAMAN"),
]

# Display order of classes in the dropdown.
CLASS_ORDER = ["DEATHKNIGHT", "DRUID", "HUNTER", "MAGE", "PALADIN", "PRIEST", "ROGUE", "SHAMAN", "WARLOCK", "WARRIOR"]

# The workbook's slot spellings, normalised to what the addon shows.
SLOT_NAMES = {
    "leggs": "Legs",
    "ring": "Ring",
    "trinket": "Trinket",
}


def class_for_sheet(name):
    for keyword, token in CLASS_KEYWORDS:
        if keyword.lower() in name.lower():
            return token
    return None


def spec_key(name):
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def lua_string(text):
    return '"' + text.replace("\\", "\\\\").replace('"', '\\"') + '"'


def read_roles(wb):
    """Map sheet name -> role label from the navigation sheet's hyperlinks."""
    roles = {}
    if NAV_SHEET not in wb.sheetnames:
        return roles
    ws = wb[NAV_SHEET]
    # Role headers sit on one row; each spec link below belongs to the header
    # whose column block it is in. Blocks are 3 columns wide (header, link, gap).
    headers = {}
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip().isupper() and not cell.hyperlink:
                headers[cell.column] = cell.value.strip().title()
    if not headers:
        return roles
    header_cols = sorted(headers)
    for row in ws.iter_rows():
        for cell in row:
            link = cell.hyperlink
            if not link or not link.location:
                continue
            target = link.location.split("!")[0].strip("'")
            block = max((c for c in header_cols if c <= cell.column), default=header_cols[0])
            roles[target] = headers[block]
    return roles


def read_item(cell, row_cells):
    """Return (ilvl, item_id, name) for an item cell, or None."""
    if not isinstance(cell.value, str):
        return None
    name = cell.value.strip()
    if not name:
        return None
    item_id = None
    if cell.hyperlink and cell.hyperlink.target:
        match = ITEM_LINK.search(cell.hyperlink.target)
        if match:
            item_id = int(match.group(1))
    # Item names in this workbook are written with a leading space; a plain
    # string without a link and without that marker is a header, not an item.
    if item_id is None and not cell.value.startswith(" "):
        return None
    ilvl = None
    for other in reversed(row_cells[: cell.column - 1]):
        if isinstance(other.value, (int, float)):
            ilvl = int(other.value)
            break
        if other.value not in (None, ""):
            break
    return ilvl, item_id, name


def read_spec(ws):
    items = []
    for row in ws.iter_rows():
        slot_cell = row[0]
        slot = slot_cell.value
        if not isinstance(slot, str) or not slot.strip():
            continue
        slot = slot.strip()
        slot = SLOT_NAMES.get(slot.lower(), slot)
        found = []
        for cell in row[1:]:
            item = read_item(cell, row)
            if item:
                found.append(item)
        if not found:
            continue
        ilvl, item_id, name = found[0]
        entry = {"slot": slot, "ilvl": ilvl, "id": item_id, "name": name, "alts": []}
        for alt_ilvl, alt_id, alt_name in found[1:]:
            entry["alts"].append({"ilvl": alt_ilvl, "id": alt_id, "name": alt_name})
        items.append(entry)
    return items


def emit_item(out, item, indent):
    pad = " " * indent
    fields = ["slot = " + lua_string(item["slot"])]
    if item["ilvl"] is not None:
        fields.append("ilvl = %d" % item["ilvl"])
    if item["id"] is not None:
        fields.append("id = %d" % item["id"])
    fields.append("name = " + lua_string(item["name"]))
    if item.get("alts"):
        out.append(pad + "{ " + ", ".join(fields) + ",")
        out.append(pad + "    alts = {")
        for alt in item["alts"]:
            alt_fields = []
            if alt["ilvl"] is not None:
                alt_fields.append("ilvl = %d" % alt["ilvl"])
            if alt["id"] is not None:
                alt_fields.append("id = %d" % alt["id"])
            alt_fields.append("name = " + lua_string(alt["name"]))
            out.append(pad + "        { " + ", ".join(alt_fields) + " },")
        out.append(pad + "    },")
        out.append(pad + "},")
    else:
        out.append(pad + "{ " + ", ".join(fields) + " },")


def build(workbook_path, output_path):
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    roles = read_roles(wb)

    specs = []
    for name in wb.sheetnames:
        if name == NAV_SHEET:
            continue
        class_token = class_for_sheet(name)
        if not class_token:
            print("skipping sheet with unknown class: %s" % name, file=sys.stderr)
            continue
        items = read_spec(wb[name])
        if not items:
            print("skipping sheet with no items: %s" % name, file=sys.stderr)
            continue
        specs.append({
            "key": spec_key(name),
            "label": name,
            "class": class_token,
            "role": roles.get(name, ""),
            "items": items,
        })

    # Stable order: class order, then workbook order within a class.
    order = {token: i for i, token in enumerate(CLASS_ORDER)}
    specs.sort(key=lambda s: (order.get(s["class"], 99), wb.sheetnames.index(s["label"])))

    total_items = sum(len(s["items"]) for s in specs)
    total_alts = sum(len(i["alts"]) for s in specs for i in s["items"])
    missing_ids = [(s["label"], i["name"]) for s in specs for i in s["items"] if i["id"] is None]
    missing_ids += [(s["label"], a["name"]) for s in specs for i in s["items"] for a in i["alts"] if a["id"] is None]

    out = []
    out.append("-- GENERATED FILE - do not edit by hand.")
    out.append("-- Built by tools/build_bis_data.py from:")
    out.append("--   " + os.path.basename(workbook_path))
    out.append("-- Regenerate with:  python tools/build_bis_data.py")
    out.append("--")
    out.append("-- %d specs, %d best-in-slot items, %d alternatives." % (len(specs), total_items, total_alts))
    if missing_ids:
        out.append("-- Entries without an item id (no link in the workbook) are resolved by")
        out.append("-- name through AtlasLoot at runtime:")
        for label, item_name in missing_ids:
            out.append("--   %s: %s" % (label, item_name))
    out.append("")
    out.append("RaidInspectorBiSData = {")
    out.append("    source = " + lua_string(os.path.basename(workbook_path)) + ",")
    out.append("    specs = {")
    for spec in specs:
        out.append("        {")
        out.append("            key = " + lua_string(spec["key"]) + ",")
        out.append("            label = " + lua_string(spec["label"]) + ",")
        out.append("            class = " + lua_string(spec["class"]) + ",")
        out.append("            role = " + lua_string(spec["role"]) + ",")
        out.append("            items = {")
        for item in spec["items"]:
            emit_item(out, item, 16)
        out.append("            },")
        out.append("        },")
    out.append("    },")
    out.append("}")
    out.append("")

    with open(output_path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(out))

    print("wrote %s: %d specs, %d items, %d alternatives, %d without id"
          % (output_path, len(specs), total_items, total_alts, len(missing_ids)))
    for label, item_name in missing_ids:
        print("  no id: %s / %s" % (label, item_name))


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbook", nargs="?", default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", default=os.path.join(here, "..", "RaidInspector_BiS.lua"))
    args = parser.parse_args()
    if not os.path.isfile(args.workbook):
        sys.exit("workbook not found: %s" % args.workbook)
    build(args.workbook, os.path.normpath(args.output))


if __name__ == "__main__":
    main()
